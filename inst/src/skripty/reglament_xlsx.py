#!/usr/bin/env python3
"""Регламент в виде книги xlsx (для Google Таблиц): выгрузка из репозитория и обратное чтение.

    python3 inst/src/skripty/reglament_xlsx.py --out ФАЙЛ.xlsx   выгрузить: шаги, обязанности, должности, места…
    python3 inst/src/skripty/reglament_xlsx.py --check           выгрузить во временный файл, прочитать, сверить
    python3 inst/src/skripty/reglament_xlsx.py --diff ФАЙЛ.xlsx  чем правленая книга отличается от репозитория

Каждый лист — как листы таблицы проекта: служебные строки с метками в столбце C (header_first, key, вид столбца,
var_key, var_name, var_value, hide_col, header_08 … header_last), под header_last — строка заголовков, ниже данные.
Строки header_08, info_fill, header_last — вики: что это, как заполнять, где видно на схеме.
Столбцы находятся по служебному имени в строке 2 (key), а не по букве: их можно переставлять.
Объединённых ячеек нет. Один шаг — одна или несколько строк: опознающие столбцы повторяются, значения из списков
(места, кто согласует, статусы…) — по строке на значение; остальное — в первой строке шага.
Нужен openpyxl (apt: python3-openpyxl); для make test не нужен.
"""
import os
import re
import sys
import tempfile

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import protsessy  # noqa: E402
import sobrat  # noqa: E402
import yaml  # noqa: E402

TECH = ['header_first', 'key', None, 'var_key', 'var_name', 'var_value', 'hide_col', 'header_08', 'info_fill', 'header_last']
R_KEY, R_KIND, R_VNAME, R_VVAL, R_HIDE, R_CHTO, R_KAK, R_SKHEMA = 2, 3, 5, 6, 7, 8, 9, 10
R_HEAD = len(TECH) + 1          # строка заголовков
R_DATA = R_HEAD + 1             # первая строка данных
SPARE = 300                     # сколько пустых строк под новые данные получают списки и проверки
README = 'Как читать'
OBJECTS = {'proekt': 'Проект', 'sdelka': 'Сделка', 'tovar': 'Товар'}
MARKS = ('И', 'С', 'Сг', 'У', 'В')
MARK_TEXT = {'И': 'исполнитель', 'С': 'соисполнитель', 'Сг': 'соисполнитель: Согласование(консультант)',
             'У': 'соисполнитель: Информируемый', 'В': 'соисполнитель: Владелец процесса'}
GREEN, YELLOW = 'FF00FF00', 'FFFFFF00'   # как в регламенте: исполнитель — зелёный, остальные отметки — жёлтый
LAYER = {'id': 'FFE7E6E6', 'reg': 'FFFFF2CC', 'tolk': 'FFDDEBF7', 'svc': 'FFF2F2F2', 'spr': 'FFE2EFDA'}
BAND = ('FFFFFFFF', 'FFF5F5F5')
README_URL_ROW = {'reglament': 20, 'papki': 21}   # ячейки C20, C21 листа «Как читать» — адреса оригинала в Google


def col(spec_key, title, chto, kak='', skhema='', width=14, layer='tolk', kind='', many=False, ident=False,
        vname='', vval='', hide='', check=None, valid=None, wrap=True):
    """Столбец листа. wrap=False — длинный текст в одну строку, чтобы шаг не растягивался по высоте."""
    return dict(key=spec_key, title=title, chto=chto, kak=kak, skhema=skhema, width=width, layer=layer, kind=kind,
                many=many, ident=ident, vname=vname, vval=vval, hide=hide, check=check, valid=valid, wrap=wrap)


def rng(sheet, letter):
    return f"'{sheet}'!${letter}${R_DATA}:${letter}${R_DATA + 2000}"


# ---------- данные из репозитория ----------

def repo():
    parts, reg, _, errors = sobrat.load()
    if errors:
        sys.exit('FAIL: ' + '; '.join(errors))
    procs = protsessy.procs(parts)
    papki = yaml.safe_load(open(f'{sobrat.SRC}/md-kod/pro/papki.yaml', encoding='utf-8')) or []
    return parts, reg, procs, papki


def title_of(p):
    return re.sub(r'^Процесс\.\s*', '', p['title'])


def statuses(procs, reg):
    """Объект → статусы по порядку: из шагов и из описания статусов проекта в Битрикс24."""
    out = {o: [] for o in OBJECTS}
    for p in procs:
        for s in p['shagi']:
            for o, st in (p['tolk'][s['id']].get('status') or {}).items():
                if st not in out[o]:
                    out[o].append(st)
    known = re.findall(r'«([^»]+)»', reg['mes'].get('bitriks-status-proekta', {}).get('polno', ''))
    out['proekt'] = list(dict.fromkeys(out['proekt'] + known))
    return out


# ---------- описание листов ----------

def spec_reglament():
    O, R, M = 'Обязанности', 'Должности', 'Места'
    cols = [
        col('err', 'Проверка', 'Что не так в строке: пусто — всё в порядке.', 'Не заполнять: считает формула.', '—',
            width=16, layer='svc', kind='formulaCell'),
        col('Оригинал', 'Оригинал', 'Ссылка на строку исходного регламента.', 'Не заполнять: формула по «Ориг. строка» и адресу на листе «Как читать».',
            '—', width=9, layer='svc', kind='formulaCell'),
        col('IdСтроки', 'ID_Строки', 'Ключ строки: ID шага и номер строки внутри шага.', 'Не заполнять: формула.', '—',
            width=10, layer='svc', kind='formulaCell'),
        col('Процесс', 'Процесс', 'Процесс регламента.', 'Выбрать из списка (лист «Процессы»). Повторять во всех строках шага.',
            'страница процесса', width=18, layer='id', ident=True, valid=rng('Процессы', 'D')),
        col('IdШага', 'ID шага', 'Постоянный ключ шага: код процесса и номер (zk-16). По нему шаг связан со схемами и репозиторием.',
            'Новый шаг — код процесса и номер; без номера — x1, x2. Повторять во всех строках шага.', 'узел на схеме',
            width=8, layer='id', ident=True),
        col('НомерШага', '№', 'Номер шага в регламенте — текстом (иначе «1.1» станет датой).', 'Повторять во всех строках шага.',
            'номер в подписи узла', width=5, layer='id', ident=True),
        col('Шаг', 'Шаг', 'Что делается — как в регламенте.', 'Повторять во всех строках шага (в продолжении — серым).',
            'подпись узла', width=45, layer='id', ident=True),
        col('Раздел', 'Раздел', 'Подраздел процесса: «Доплаты», «ПРИЕМКА»…', 'Повторять во всех строках шага.', '—',
            width=14, layer='id', ident=True),
    ]
    for c, name in protsessy.COLUMNS.items():
        cols.append(col(f'Рег{c}', protsessy.col_name(c), f'Столбец {c}', 'И, С, Сг, У, В', 'дорожка',
                        width=7, layer='reg', vname='Отметки регламента', vval='; '.join(MARKS), valid='"' + ','.join(MARKS) + '"'))
    cols += [
        col('РегПояснения', 'Отметки: пояснения', 'Текст в ячейках отметок регламента, если он не просто «исполнитель» или «согласование».',
            'Формат «столбец: текст», через «;».', '—', width=24, layer='reg', wrap=False),
        col('РегПримечание', 'Примечание по шагу', 'Столбец «Примечание по шагу» регламента.', 'Только в первой строке шага.', '—',
            width=30, layer='reg', hide='Сгрупировать', wrap=False),
        col('РегВСистеме', 'Действия в таблице/Битриксе', 'Столбец «Действия в таблице/Битриксе» регламента.', 'Только в первой строке шага.',
            '—', width=30, layer='reg', hide='Сгрупировать', wrap=False),
        col('РегИзменения', 'Изменения в системе', 'Столбец «ИЗМЕНЕНИЯ В СИСТЕМЕ» регламента.', 'Только в первой строке шага.', '—',
            width=20, layer='reg', hide='Сгрупировать', wrap=False),
        col('ОригСтрока', 'Ориг. строка', 'Номер строки шага в исходном регламенте (лист «10.10.2026 Версия К»).',
            'Не менять. У нового шага — пусто.', '—', width=8, layer='reg'),
        col('Ставит', 'Ставит задачу', 'Обязанность, которая ставит задачу на этот шаг.', 'Выбрать из списка (лист «Обязанности»). Несколько — по строке на каждую.',
            'стрелка «ставит» в схеме «По обязанностям»', width=24, many=True, valid=rng(O, 'D')),
        col('Делает', 'Делает', 'Обязанность, которая делает шаг. Ровно одна.', 'Выбрать из списка; только в первой строке шага.',
            'узел обязанности; дорожка — должность из «Предполагаем»', width=24, valid=rng(O, 'D')),
        col('Принимает', 'Принимает результат', 'Обязанность, которой передаётся результат шага.', 'Из списка; несколько — по строке на каждую.',
            'стрелка «передаёт»', width=24, many=True, valid=rng(O, 'D')),
        col('Передаёт', 'Передаёт (что)', 'Что шаг отдаёт дальше: задача, счёт, лист, документ, товар.',
            'Из списка (лист «Что передаётся»); несколько — по строке.', 'подпись стрелки на дорожках', width=24, many=True,
            valid=rng('Что передаётся', 'D')),
        col('Отслеживает', 'Отслеживает', 'Обязанность, которая следит за шагом.', 'Из списка; несколько — по строке.', 'пунктир «следит»',
            width=20, many=True, valid=rng(O, 'D')),
        col('Согласуют', 'Согласуют', 'Обязанности, с которыми шаг согласуют.', 'Из списка; несколько — по строке.', 'пунктир «согласует»',
            width=22, many=True, valid=rng(O, 'D')),
        col('Помогают', 'Помогают (должности)', 'Должности-соисполнители.', 'Из списка (лист «Должности»); несколько — по строке.',
            'таблица шагов', width=20, many=True, valid=rng(R, 'D')),
        col('Узнают', 'Узнают о результате (должности)', 'Должности, которых уведомляют.', 'Из списка (лист «Должности»); несколько — по строке.',
            'таблица шагов', width=20, many=True, valid=rng(R, 'D')),
        col('Место', 'Где', 'Место работы: лист, задача Битрикс24, 1С, папка, почта, ЭДО.', 'Из списка (лист «Места»); несколько — по строке.',
            'схема «Где работают»', width=24, many=True, valid=rng(M, 'D')),
        col('ОбъектСтатуса', 'Объект статуса', 'Чей статус меняет шаг.', 'Проект, Сделка или Товар; пара с «Статус после» — в одной строке.',
            'схема «Статусы»', width=10, many=True, vname='Объекты статусов', vval='; '.join(OBJECTS.values()),
            valid='"' + ','.join(OBJECTS.values()) + '"'),
        col('СтатусПосле', 'Статус после', 'Какой статус ставит шаг.', 'Из списка (лист «Статусы»); пара с «Объект статуса».',
            'переход на схеме «Статусы»', width=20, many=True, valid=rng('Статусы', 'E')),
        col('Догадка', 'Догадка', 'Что в толковании шага не подтверждено.', 'Только в первой строке шага.', 'курсив в таблице шагов',
            width=30, wrap=False),
    ]
    return cols


def spec_dict(kind):
    """Справочники: обязанности, должности, места, процессы, статусы, папки."""
    common = [col('err', 'Проверка', 'Что не так в строке: пусто — всё в порядке.', 'Не заполнять: формула.', '—', width=14,
                  layer='svc', kind='formulaCell')]
    st_valid = '"готово,проверить,заглушка"'
    if kind == 'Обязанности':
        return common + [
            col('Шагов', 'Шагов', 'Сколько шагов делает обязанность (лист «Регламент», «Делает»).', 'Не заполнять: формула.', '—',
                width=7, layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Постоянный ключ (латиница): по нему обязанность связана с репозиторием и карточкой.',
                'Новой — можно не заполнять: сборщик сделает из названия. Повторять во всех строках.', 'имя карточки', width=22, layer='id', ident=True),
            col('Название', 'Название', '«Тот, кто …»: одна обязанность — один ответственный.', 'Уникальное; его выбирают в «Регламенте». Повторять во всех строках.',
                'узел на схемах', width=32, layer='id', ident=True),
            col('ЗаЧтоОтвечает', 'За что отвечает', 'Одна фраза.', 'Только в первой строке.', 'карточка', width=45, layer='spr'),
            col('Предполагаем', 'Предполагаем (должность)', 'Какая должность, по догадке, выполняет. Пусто — не знаем.',
                'Из списка (лист «Должности»); несколько — по строке (это вопрос: ответственный должен быть один).',
                'дорожка на схеме «По обязанностям»', width=26, many=True, layer='spr', valid=rng('Должности', 'D')),
            col('Контур', 'Контур', 'Контур, где обязанность работает.', 'Только в первой строке.', 'карточка должности: группы по контурам',
                width=10, layer='spr'),
            col('Где', 'Где (без шагов)', 'Места работы, если шагов в регламенте нет (иначе места берутся из шагов).',
                'Из списка (лист «Места»); несколько — по строке.', 'карточка', width=22, many=True, layer='spr', valid=rng('Места', 'D')),
            col('Источник', 'Источник', 'Откуда взято: регламент, роли.docx, смарт-процесс…', 'Через «; ».', 'карточка', width=18, layer='spr'),
            col('Почему', 'Почему выделили', 'Почему это отдельная обязанность.', '', 'карточка', width=40, layer='spr'),
            col('Догадка', 'Догадка', 'Что не подтверждено.', '', 'карточка; список расхождений', width=40, layer='spr'),
            col('Статус', 'Статус', 'готово, проверить или заглушка.', 'Из списка.', '«черновик» в карточке', width=10, layer='spr',
                vname='Статусы записи', vval='готово; проверить; заглушка', valid=st_valid),
        ]
    if kind == 'Должности':
        return common + [
            col('Обязанностей', 'Обязанностей', 'Сколько обязанностей ей приписано (лист «Обязанности»).', 'Не заполнять: формула.', '—',
                width=8, layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Постоянный ключ (латиница).', 'Повторять во всех строках.', 'имя карточки', width=14, layer='id', ident=True),
            col('Название', 'Название', 'Должность — кто по факту.', 'Уникальное. Повторять во всех строках.', 'дорожка на схемах', width=30, layer='id', ident=True),
            col('Описание', 'Описание', 'Что делает — одной фразой.', '', 'карточка', width=45, layer='spr'),
            col('Контур', 'Контур', 'Где в основном работает.', '', 'карточка', width=16, layer='spr'),
            col('Столбец', 'Столбец регламента', 'Её столбец в регламенте (H…N): по нему сверяются регламент и догадки.',
                'Буква; несколько — по строке.', 'список расхождений', width=8, many=True, layer='spr'),
            col('НазываютТакже', 'Как ещё называют', 'Названия в других текстах.', 'Несколько — по строке.', 'карточка', width=30,
                many=True, layer='spr'),
            col('Источник', 'Источник', 'Откуда взято.', 'Через «; ».', 'карточка', width=18, layer='spr'),
            col('Почему', 'Почему выделили', '', '', 'карточка', width=40, layer='spr'),
            col('Догадка', 'Догадка', 'Что не подтверждено.', '', 'карточка', width=40, layer='spr'),
            col('Статус', 'Статус', 'готово, проверить или заглушка.', 'Из списка.', '—', width=10, layer='spr',
                vname='Статусы записи', vval='готово; проверить; заглушка', valid=st_valid),
        ]
    if kind == 'Места':
        vids = ['лист', 'таблица', 'битрикс', '1с', 'диск', 'почта', 'эдо', 'внешнее']
        return common + [
            col('Шагов', 'Шагов', 'Сколько строк «Регламента» ссылаются на место.', 'Не заполнять: формула.', '—', width=7,
                layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Постоянный ключ (латиница).', '', '—', width=22, layer='id', ident=True),
            col('Название', 'Название', 'Лист, таблица, Битрикс24, 1С, папка…', 'Уникальное.', 'узел на схеме «Где работают»', width=36, layer='id', ident=True),
            col('Описание', 'Описание', '', '', 'страница «Места работы»', width=50, layer='spr'),
            col('Вид', 'Вид', 'Вид места.', 'Из списка.', 'форма узла на схеме', width=10, layer='spr', vname='Виды мест',
                vval='; '.join(vids), valid='"' + ','.join(vids) + '"'),
            col('Часть', 'Часть с описанием', 'Часть инструкций, где место описано (например lis.tab-proekta).', 'Ключ части.', 'ссылка', width=22, layer='spr'),
        ]
    if kind == 'Что передаётся':
        vids = ['задача', 'документ', 'данные', 'товар', 'деньги', 'сообщение']
        return common + [
            col('Шагов', 'Шагов', 'Сколько строк «Регламента» передают предмет.', 'Не заполнять: формула.', '—', width=7,
                layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Постоянный ключ (латиница).', '', '—', width=22, layer='id', ident=True),
            col('Название', 'Название', 'Что передаётся.', 'Уникальное.', 'подпись стрелки на дорожках', width=40, layer='id', ident=True),
            col('Вид', 'Вид', 'Вид предмета.', 'Из списка.', '—', width=10, layer='spr', vname='Виды предметов',
                vval='; '.join(vids), valid='"' + ','.join(vids) + '"'),
            col('Где', 'Где лежит', 'Место, где предмет хранится или виден.', 'Из списка (лист «Места»).', 'страница «Что передаётся»',
                width=30, layer='spr', valid=rng('Места', 'D')),
            col('Догадка', 'Догадка', 'Что не подтверждено.', '', '—', width=40, layer='spr'),
        ]
    if kind == 'Процессы':
        return common + [
            col('Шагов', 'Шагов', 'Сколько шагов в процессе (лист «Регламент»).', 'Не заполнять: формула.', '—', width=7,
                layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Папка процесса в репозитории.', 'Повторять во всех строках.', 'адрес страницы', width=30, layer='id', ident=True),
            col('Название', 'Название', 'Как процесс называется у нас.', 'Уникальное. Повторять во всех строках.', 'заголовок страницы', width=34, layer='id', ident=True),
            col('КодШагов', 'Код шагов', 'Начало ID шагов процесса (zk → zk-16).', '', '—', width=8, layer='spr'),
            col('ВРегламенте', 'В регламенте', 'Как процесс назван в исходном регламенте.', '', '—', width=34, layer='spr'),
            col('Контур', 'Контур', 'Контур процесса.', '', 'карта процессов: группа', width=8, layer='spr'),
            col('Дальше', 'Дальше', 'Какой процесс идёт следом.', 'Из списка; несколько — по строке.', 'стрелка на карте процессов',
                width=30, many=True, layer='spr', valid=rng('Процессы', 'D')),
        ]
    if kind == 'Статусы':
        return common + [
            col('Шагов', 'Шагов', 'Сколько шагов ставит статус.', 'Не заполнять: формула.', '—', width=7, layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Ключ строки.', '', '—', width=10, layer='id', ident=True),
            col('Объект', 'Объект', 'Чей статус.', 'Проект, Сделка или Товар.', 'отдельная схема на объект', width=10, layer='spr',
                valid='"' + ','.join(OBJECTS.values()) + '"'),
            col('Статус', 'Статус', 'Значение статуса — как в системе.', 'Уникально внутри объекта.', 'состояние на схеме', width=34, layer='spr'),
        ]
    if kind == 'Папки':
        return common + [
            col('Оригинал', 'Оригинал', 'Ссылка на строку листа «Структура папок».', 'Не заполнять: формула.', '—', width=9,
                layer='svc', kind='formulaCell'),
            col('Ключ', 'Ключ', 'Ключ строки: p + номер строки оригинала.', '', '—', width=7, layer='id', ident=True),
            col('Уровень', 'Уровень', 'Уровень вложенности папки (1 — папка проекта).', '', 'отступ', width=8, layer='spr'),
            col('Папка', 'Папка', 'Название папки.', '', '—', width=40, layer='spr'),
            col('Создаёт', 'Кто создаёт', 'Кто создаёт папку — как в оригинале.', '', '—', width=26, layer='spr'),
            col('Наполняет', 'Кто наполняет', 'Кто кладёт документы.', '', '—', width=26, layer='spr'),
            col('Примечание', 'Примечание', '', '', '—', width=40, layer='spr'),
            col('ОригСтрока', 'Ориг. строка', 'Строка в листе «Структура папок» исходного файла.', 'Не менять.', '—', width=8, layer='spr'),
        ]
    raise KeyError(kind)


# ---------- строки данных ----------

def rows_reglament(procs, reg):
    oby, rol, mes = reg['oby'], reg['rol'], reg['mes']
    out = []   # (номер шага по порядку, {ключ столбца: значение}) — строка; номер для чередования цвета
    for p in procs:
        for s in p['shagi']:
            t = p['tolk'][s['id']]
            marks, notes = {}, []
            for c, v in (s.get('kolonki') or {}).items():
                m = protsessy.mark(v)
                marks[c] = m
                if v != MARK_TEXT.get(m) and v not in ('исполнитель: Исполнитель', 'соисполнитель: Соисполнитель',
                                                       'соисполнитель: Соисполнители'):
                    notes.append(f'{c}: {v}')
            first = {'Процесс': title_of(p), 'IdШага': s['id'], 'НомерШага': s['nomer'], 'Шаг': s['title'],
                     'Раздел': s.get('razdel', ''), 'РегПояснения': '; '.join(notes), 'РегПримечание': s.get('primechanie', ''),
                     'РегВСистеме': s.get('v_sisteme', ''), 'РегИзменения': s.get('izmeneniya', ''), 'ОригСтрока': s.get('stroka', ''),
                     'Делает': oby[t['delaet']]['title'], 'Догадка': t.get('dogadka', '')}
            first.update({f'Рег{c}': m for c, m in marks.items()})
            lists = {
                'Ставит': [oby[v]['title'] for v in protsessy.as_list(t.get('stavit'))],
                'Принимает': [oby[v]['title'] for v in protsessy.as_list(t.get('prinimaet'))],
                'Передаёт': [reg['art'][v]['title'] for v in protsessy.as_list(t.get('peredaet'))],
                'Отслеживает': [oby[v]['title'] for v in protsessy.as_list(t.get('otslezhivaet'))],
                'Согласуют': [oby[v]['title'] for v in protsessy.as_list(t.get('soglasuyut'))],
                'Помогают': [rol[v[4:]]['title'] for v in protsessy.as_list(t.get('pomogayut'))],
                'Узнают': [rol[v[4:]]['title'] for v in protsessy.as_list(t.get('uznayut'))],
                'Место': [mes[v]['title'] for v in protsessy.as_list(t.get('mesto'))],
                'ОбъектСтатуса': [OBJECTS[o] for o in (t.get('status') or {})],
                'СтатусПосле': list((t.get('status') or {}).values()),
            }
            n = max([1] + [len(v) for v in lists.values()])
            for i in range(n):
                row = dict(first) if i == 0 else {k: first[k] for k in ('Процесс', 'IdШага', 'НомерШага', 'Шаг', 'Раздел')}
                row.update({k: v[i] for k, v in lists.items() if i < len(v)})
                out.append((s['id'], row))
    return out


def multi(first, lists, ident):
    n = max([1] + [len(v) for v in lists.values()])
    rows = []
    for i in range(n):
        row = dict(first) if i == 0 else {k: first[k] for k in ident}
        row.update({k: v[i] for k, v in lists.items() if i < len(v)})
        rows.append(row)
    return rows


def rows_dict(kind, reg, procs, papki, parts):
    out = []
    if kind == 'Обязанности':
        for k, o in reg['oby'].items():
            first = {'Ключ': k, 'Название': o['title'], 'ЗаЧтоОтвечает': o.get('polno', ''), 'Контур': str(o.get('kontur', '')),
                     'Источник': '; '.join(protsessy.as_list(o.get('istochnik'))), 'Почему': o.get('pochemu', ''),
                     'Догадка': o.get('dogadka', ''), 'Статус': o.get('status', '')}
            lists = {'Предполагаем': [reg['rol'][r]['title'] for r in protsessy.as_list(o.get('predpolagaem'))],
                     'Где': [reg['mes'][m]['title'] for m in protsessy.as_list(o.get('gde'))]}
            out += [(k, r) for r in multi(first, lists, ('Ключ', 'Название'))]
    elif kind == 'Должности':
        for k, r in reg['rol'].items():
            first = {'Ключ': k, 'Название': r['title'], 'Описание': r.get('polno', ''), 'Контур': str(r.get('kontur', '')),
                     'Источник': '; '.join(protsessy.as_list(r.get('istochnik'))), 'Почему': r.get('pochemu', ''),
                     'Догадка': r.get('dogadka', ''), 'Статус': r.get('status', '')}
            lists = {'Столбец': protsessy.as_list(r.get('kolonka')), 'НазываютТакже': protsessy.as_list(r.get('nazvaniya'))}
            out += [(k, x) for x in multi(first, lists, ('Ключ', 'Название'))]
    elif kind == 'Что передаётся':
        for k, a in reg['art'].items():
            out.append((k, {'Ключ': k, 'Название': a['title'], 'Вид': a.get('vid', ''),
                            'Где': reg['mes'][a['gde']]['title'] if a.get('gde') else '', 'Догадка': a.get('dogadka', '')}))
    elif kind == 'Места':
        for k, m in reg['mes'].items():
            out.append((k, {'Ключ': k, 'Название': m['title'], 'Описание': m.get('polno', ''), 'Вид': m.get('vid', ''),
                            'Часть': m.get('chast', '')}))
    elif kind == 'Процессы':
        for p in procs:
            first = {'Ключ': p['slug'], 'Название': title_of(p), 'КодШагов': p['shagi'][0]['id'].split('-')[0],
                     'ВРегламенте': p['reglament_title'], 'Контур': p['kontur']}
            lists = {'Дальше': [title_of(parts['pro.' + n]) for n in p['dalee']]}
            out += [(p['slug'], r) for r in multi(first, lists, ('Ключ', 'Название'))]
    elif kind == 'Статусы':
        i = 0
        for o, sts in statuses(procs, reg).items():
            for st in sts:
                i += 1
                out.append((f's{i}', {'Ключ': f'{o}-{i}', 'Объект': OBJECTS[o], 'Статус': st}))
    elif kind == 'Папки':
        for f in papki:
            k = f'p{f["stroka"]}'
            out.append((k, {'Ключ': k, 'Уровень': f['uroven'], 'Папка': f['papka'], 'Создаёт': f.get('sozdaet', ''),
                            'Наполняет': f.get('napolnyaet', ''), 'Примечание': f.get('primechanie', ''), 'ОригСтрока': f['stroka']}))
    return out


# ---------- формулы ----------

def formulas(kind, spec, r):
    """Формулы служебных столбцов строки r: одинаково работают в Excel и Google Таблицах."""
    L = {c['key']: get_column_letter(i + 1) for i, c in enumerate(spec)}
    f = {}
    if kind == 'Регламент':
        id_, first = f'${L["IdШага"]}{r}', f'COUNTIF(${L["IdШага"]}${R_DATA}:${L["IdШага"]}{r},${L["IdШага"]}{r})=1'
        f['IdСтроки'] = f'=IF({id_}="","",{id_}&"."&COUNTIF(${L["IdШага"]}${R_DATA}:${L["IdШага"]}{r},{id_}))'
        url = f"'{README}'!$C${README_URL_ROW['reglament']}"
        orig = f'${L["ОригСтрока"]}{r}'
        f['Оригинал'] = f'=IF({orig}="","",IF({url}="","стр. "&{orig},HYPERLINK({url}&"&range=A"&{orig},"стр. "&{orig})))'
        checks = [f'IF(AND({first},${L["Делает"]}{r}=""),"нет «Делает»; ","")']
        for key, sheet in (('Ставит', 'Обязанности'), ('Делает', 'Обязанности'), ('Принимает', 'Обязанности'), ('Передаёт', 'Что передаётся'),
                           ('Отслеживает', 'Обязанности'), ('Согласуют', 'Обязанности'), ('Помогают', 'Должности'),
                           ('Узнают', 'Должности'), ('Место', 'Места'), ('СтатусПосле', 'Статусы')):
            letter = 'E' if sheet == 'Статусы' else 'D'
            cell = f'${L[key]}{r}'
            checks.append(f'IF(AND({cell}<>"",COUNTIF({rng(sheet, letter)},{cell})=0),"{key}: нет в «{sheet}»; ","")')
        checks.append(f'IF((${L["ОбъектСтатуса"]}{r}="")<>(${L["СтатусПосле"]}{r}=""),"статус без объекта; ","")')
        f['err'] = f'=IF({id_}="","",' + '&'.join(checks) + ')'
    elif kind == 'Обязанности':
        name = f'${L["Название"]}{r}'
        f['Шагов'] = f"=IF({name}=\"\",\"\",COUNTIF({rng('Регламент', reg_letter('Делает'))},{name}))"
        f['err'] = (f'=IF({name}="","",IF(COUNTIF(${L["Название"]}${R_DATA}:${L["Название"]}{r},{name})>1,'
                    f'IF(${L["Ключ"]}{r}<>INDEX(${L["Ключ"]}${R_DATA}:${L["Ключ"]}{r},MATCH({name},${L["Название"]}${R_DATA}:${L["Название"]}{r},0)),"название уже занято; ",""),"")'
                    f'&IF(AND(${L["Предполагаем"]}{r}<>"",COUNTIF({rng("Должности", "D")},${L["Предполагаем"]}{r})=0),"должности нет; ","")'
                    f'&IF(AND(${L["Где"]}{r}<>"",COUNTIF({rng("Места", "D")},${L["Где"]}{r})=0),"места нет; ",""))')
    elif kind == 'Должности':
        name = f'${L["Название"]}{r}'
        f['Обязанностей'] = f"=IF({name}=\"\",\"\",COUNTIF({rng('Обязанности', dict_letter('Обязанности', 'Предполагаем'))},{name}))"
        f['err'] = ''
    elif kind == 'Места':
        name = f'${L["Название"]}{r}'
        f['Шагов'] = f"=IF({name}=\"\",\"\",COUNTIF({rng('Регламент', reg_letter('Место'))},{name}))"
        f['err'] = ''
    elif kind == 'Что передаётся':
        name = f'${L["Название"]}{r}'
        f['Шагов'] = f"=IF({name}=\"\",\"\",COUNTIF({rng('Регламент', reg_letter('Передаёт'))},{name}))"
        f['err'] = f'=IF(AND(${L["Где"]}{r}<>"",COUNTIF({rng("Места", "D")},${L["Где"]}{r})=0),"места нет","")'
    elif kind == 'Процессы':
        name = f'${L["Название"]}{r}'
        f['Шагов'] = (f"=IF({name}=\"\",\"\",COUNTIFS({rng('Регламент', reg_letter('Процесс'))},{name},"
                      f"{rng('Регламент', reg_letter('IdСтроки'))},\"*.1\"))")
        f['err'] = ''
    elif kind == 'Статусы':
        st = f'${L["Статус"]}{r}'
        f['Шагов'] = f"=IF({st}=\"\",\"\",COUNTIF({rng('Регламент', reg_letter('СтатусПосле'))},{st}))"
        f['err'] = ''
    elif kind == 'Папки':
        url = f"'{README}'!$C${README_URL_ROW['papki']}"
        orig = f'${L["ОригСтрока"]}{r}'
        f['Оригинал'] = f'=IF({orig}="","",IF({url}="","стр. "&{orig},HYPERLINK({url}&"&range=A"&{orig},"стр. "&{orig})))'
        f['err'] = ''
    if kind in ('Обязанности', 'Должности', 'Места', 'Процессы') and 'Ключ' in L:   # счётчик — только в первой строке записи
        k = f'${L["Ключ"]}{r}'
        once = f'COUNTIF(${L["Ключ"]}${R_DATA}:${L["Ключ"]}{r},{k})>1'
        for name in ('Шагов', 'Обязанностей'):
            if name in f:
                f[name] = f'=IF({once},"",{f[name][1:]})'
    return {k: v for k, v in f.items() if v}


def reg_letter(key):
    return get_column_letter([c['key'] for c in spec_reglament()].index(key) + 1)


def dict_letter(kind, key):
    return get_column_letter([c['key'] for c in spec_dict(kind)].index(key) + 1)


# ---------- запись ----------

def frame(ws, spec, title):
    """Служебные строки, вики-строки и заголовки — по образцу листов таблицы проекта."""
    small, gray, bold = Font(size=8, color='FF808080'), Font(size=9, italic=True, color='FF595959'), Font(bold=True)
    for i, label in enumerate(TECH, 1):
        if label:
            ws.cell(i, 3, label)
    ws.cell(R_KIND, 3, 'formulaCell' if spec[2]['kind'] == 'formulaCell' else '')
    ws.cell(R_KEY, 3, 'key')
    prev, pending = None, None
    for i, c in enumerate(spec, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = c['width']
        fill = PatternFill('solid', fgColor=LAYER[c['layer']])
        group = {'svc': 'Служебные', 'id': 'Шаг' if title == 'Регламент' else 'Запись', 'reg': 'Как в регламенте',
                 'tolk': 'Как понимаем', 'spr': 'Справочник'}[c['layer']]
        if group != prev or pending:   # подпись группы — в первой её ячейке строки 1, но не в C1: там header_first
            if i == 3:
                pending = group
            else:
                ws.cell(1, i, pending or group).font = bold
                pending = None
            prev = group
        ws.cell(1, i).fill = fill
        if i != 3:
            ws.cell(R_KEY, i, c['key'])
            if c['kind']:
                ws.cell(R_KIND, i, c['kind'])
        if c['vname']:
            ws.cell(R_VNAME, i, c['vname'])
            ws.cell(R_VVAL, i, c['vval'])
        if c['hide']:
            ws.cell(R_HIDE, i, c['hide'])
        for r, field in ((R_CHTO, 'chto'), (R_KAK, 'kak'), (R_SKHEMA, 'skhema')):
            if i != 3:
                cell = ws.cell(r, i, c[field])
                cell.font, cell.alignment = gray, Alignment(wrap_text=True, vertical='top')
        h = ws.cell(R_HEAD, i, c['title'])
        h.font, h.fill, h.alignment = bold, fill, Alignment(wrap_text=True, vertical='top')
    for r in range(2, R_HIDE + 1):   # техника — мелким серым
        for i in range(1, len(spec) + 1):
            ws.cell(r, i).font = small
    if spec[1]['key'] in ('Оригинал', 'Шагов', 'Обязанностей'):   # подписи вики-строк — в столбце B
        for r, text in ((R_CHTO, 'Что это:'), (R_KAK, 'Как заполнять:'), (R_SKHEMA, 'На схеме:')):
            ws.cell(r, 2, text).font = Font(size=9, bold=True, color='FF595959')
    for r in range(2, R_HIDE + 1):   # строки 2–7 — техника: группой, чтобы свернуть
        ws.row_dimensions[r].outlineLevel = 1
    for r in (R_CHTO, R_KAK, R_SKHEMA):   # вики-строки — три строчки текста; больше — раздвинуть строку
        ws.row_dimensions[r].height = 36
    for i, c in enumerate(spec, 1):
        if c['hide'] == 'Сгрупировать':
            ws.column_dimensions[get_column_letter(i)].outlineLevel = 1


def fill_sheet(ws, kind, spec, rows):
    keys = [c['key'] for c in spec]
    ident = {c['key'] for c in spec if c['ident']}
    group, band = None, 1
    for n, (gkey, row) in enumerate(rows):
        r = R_DATA + n
        cont = gkey == group
        if not cont:
            group, band = gkey, 1 - band
        for i, key in enumerate(keys, 1):
            cell = ws.cell(r, i)
            v = row.get(key, '')
            if v != '' and v is not None:
                cell.value = ("'" + v) if isinstance(v, str) and v.startswith('=') else v   # текст с «=» — не формула
                if isinstance(v, str):
                    cell.number_format = '@'
            cell.alignment = Alignment(wrap_text=spec[i - 1]['wrap'], vertical='top')
            if key in ident:
                cell.fill = PatternFill('solid', fgColor=BAND[band])
                if cont:
                    cell.font = Font(color='FF9E9E9E')
        for key, formula in formulas(kind, spec, r).items():
            ws.cell(r, keys.index(key) + 1, formula).number_format = 'General'
    last = R_DATA + len(rows) + SPARE
    for r in range(R_DATA + len(rows), last):      # формулы и в запасных строках: новая строка сразу проверяется
        for key, formula in formulas(kind, spec, r).items():
            ws.cell(r, keys.index(key) + 1, formula)
    for i, c in enumerate(spec, 1):
        if c['valid']:
            L = get_column_letter(i)
            dv = DataValidation(type='list', formula1=c['valid'], allow_blank=True, showErrorMessage=True,
                                errorTitle='Нет в списке', error='Выберите значение из списка (справочник на отдельном листе).')
            dv.add(f'{L}{R_DATA}:{L}{last}')
            ws.add_data_validation(dv)
    keys_set = set(keys)   # закрепить строки до заголовков и столбцы до названия шага или записи
    target = 'Шаг' if kind == 'Регламент' else 'Название' if 'Название' in keys_set else 'Ключ'
    ws.freeze_panes = ws.cell(R_DATA, keys.index(target) + 2)
    if kind == 'Регламент':
        first, end = get_column_letter(keys.index('РегH') + 1), get_column_letter(keys.index('РегN') + 1)
        area = f'{first}{R_DATA}:{end}{last}'
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'{first}{R_DATA}="И"'], fill=PatternFill('solid', fgColor=GREEN)))
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'AND({first}{R_DATA}<>"",{first}{R_DATA}<>"И")'],
                                                        fill=PatternFill('solid', fgColor=YELLOW)))
    err = get_column_letter(keys.index('err') + 1)
    ws.conditional_formatting.add(f'{err}{R_DATA}:{err}{last}', FormulaRule(formula=[f'{err}{R_DATA}<>""'],
                                                                             font=Font(color='FFC00000', bold=True)))
    ws.auto_filter.ref = f'A{R_HEAD}:{get_column_letter(len(spec))}{R_DATA + len(rows) - 1}'


def readme(ws, procs, reg):
    lines = [
        ('Регламент: шаги, обязанности, должности', 'b'),
        ('Собрано из репозитория orb-work (inst/src/md-kod/pro, reg/*.yaml) скриптом inst/src/skripty/reglament_xlsx.py.', ''),
        ('Исходник шагов — регламент «Распределение задач Контрактация — Исполнение», лист «10.10.2026 Версия К».', ''),
        ('', ''),
        ('Слои', 'b'),
        ('Как в регламенте — отметки столбцов (И — исполнитель, С — соисполнитель, Сг — согласование, У — информируемый, В — владелец), цвета как в оригинале.', ''),
        ('Как понимаем — обязанности «тот, кто …»: одна обязанность — один ответственный. Должность — догадка на листе «Обязанности», столбец «Предполагаем».', ''),
        ('', ''),
        ('Как устроены строки', 'b'),
        ('Один шаг — одна или несколько строк. Шаг, процесс, ID и № повторяются во всех строках шага (в продолжении — серым).', ''),
        ('Значения-списки (кто согласует, места, статусы…) — по строке на значение. Остальное — только в первой строке шага.', ''),
        ('Объединённых ячеек нет. Столбцы ищутся по служебному имени в строке 2 — их можно переставлять; строки 2–7 — техника.', ''),
        ('Строки «Что это», «Как заполнять», «На схеме» — подсказки к каждому столбцу. «Проверка» (столбец A) пишет, что не так.', ''),
        ('Правки: после правок в Google — скачать как xlsx и отдать агенту; reglament_xlsx.py --diff покажет, что изменилось, и всё уйдёт в схемы.', ''),
        ('Ссылки на оригинал', 'b'),
        ('Когда исходный регламент будет в Google Таблицах — вставьте ниже адрес его листа (вида …/edit#gid=123). Столбец «Оригинал» станет ссылкой на строку.', ''),
        ('', ''),
        ('', ''),
        ('', ''),
        ('Адрес листа «10.10.2026 Версия К»:', 'k'),
        ('Адрес листа «Структура папок»:', 'k'),
        ('', ''),
        ('Числа', 'b'),
        (f'Процессов: {len(procs)}; шагов: {sum(len(p["shagi"]) for p in procs)}; обязанностей: {len(reg["oby"])}; '
         f'должностей: {len(reg["rol"])}; мест: {len(reg["mes"])}.', ''),
    ]
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width, ws.column_dimensions['C'].width = 4, 40, 90
    for i, (text, style) in enumerate(lines, 1):
        if style == 'k':
            ws.cell(i, 2, text).font = Font(bold=True)
            ws.cell(i, 3).fill = PatternFill('solid', fgColor=LAYER['reg'])
        else:
            c = ws.cell(i, 2, text)
            c.font = Font(bold=True, size=12) if style == 'b' else Font()
    assert ws.cell(README_URL_ROW['reglament'], 2).value.startswith('Адрес листа «10.10')


def export(path):
    parts, reg, procs, papki = repo()
    for kind in ('oby', 'rol', 'mes', 'art'):   # названия — то, что выбирают в списках: должны быть уникальны
        titles = [v['title'] for v in reg[kind].values()]
        dup = {t for t in titles if titles.count(t) > 1}
        if dup:
            sys.exit(f'FAIL: в reg/{kind}.yaml повторяются названия: {", ".join(sorted(dup))}')
    wb = openpyxl.Workbook()
    readme(wb.active, procs, reg)
    wb.active.title = README
    sheets = [('Регламент', spec_reglament(), rows_reglament(procs, reg))]
    sheets += [(k, spec_dict(k), rows_dict(k, reg, procs, papki, parts)) for k in ('Обязанности', 'Должности', 'Места', 'Что передаётся', 'Процессы', 'Статусы', 'Папки')]
    for kind, spec, rows in sheets:
        ws = wb.create_sheet(kind)
        frame(ws, spec, kind)
        fill_sheet(ws, kind, spec, rows)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return {k: len(r) for k, _, r in sheets}


# ---------- чтение ----------

def read_sheet(ws):
    """Строки данных листа как список словарей {служебное имя: значение}; столбцы — по строке key."""
    head = next(r for r in range(1, 40) if ws.cell(r, 3).value == 'header_last') + 1
    keyrow = next(r for r in range(1, head) if ws.cell(r, 3).value == 'key')
    kindrow = keyrow + 1
    cols = {}
    for i in range(1, ws.max_column + 1):
        k = ws.cell(keyrow, i).value if i != 3 else 'IdСтроки' if ws.title == 'Регламент' else 'Ключ'   # в C2 — метка «key»
        if k and ws.cell(kindrow, i).value != 'formulaCell':
            cols[k] = i
    rows = []
    for r in range(head + 1, ws.max_row + 1):
        row = {k: ws.cell(r, i).value for k, i in cols.items()}
        row = {k: (str(v).strip() if v is not None else '') for k, v in row.items()}
        if any(row.values()):
            rows.append(row)
    return rows


def groups(rows, key):
    out, cur = [], None
    for row in rows:
        if not row.get(key):
            continue
        if cur is None or cur[0] != row[key]:
            cur = (row[key], [])
            out.append(cur)
        cur[1].append(row)
    return out


def first(rows, k):
    return next((r[k] for r in rows if r.get(k)), '')


def many(rows, k):
    return [r[k] for r in rows if r.get(k)]


def read_book(path):
    """Книга → те же данные, что в репозитории: процессы, толкование шагов, обязанности, должности, места."""
    wb = openpyxl.load_workbook(path)
    data = {k: read_sheet(wb[k]) for k in ('Регламент', 'Обязанности', 'Должности', 'Места', 'Что передаётся', 'Процессы', 'Папки')}
    by_title = {}
    reg = {'oby': {}, 'rol': {}, 'mes': {}, 'art': {}}
    for sheet, kind in (('Должности', 'rol'), ('Места', 'mes'), ('Обязанности', 'oby'), ('Что передаётся', 'art')):
        for key, rows in groups(data[sheet], 'Ключ'):
            by_title[(kind, first(rows, 'Название'))] = key
    for key, rows in groups(data['Должности'], 'Ключ'):
        reg['rol'][key] = clean({'title': first(rows, 'Название'), 'polno': first(rows, 'Описание'), 'kontur': first(rows, 'Контур'),
                                 'kolonka': many(rows, 'Столбец'), 'nazvaniya': many(rows, 'НазываютТакже'),
                                 'istochnik': split(first(rows, 'Источник')), 'pochemu': first(rows, 'Почему'),
                                 'dogadka': first(rows, 'Догадка'), 'status': first(rows, 'Статус')})
    for key, rows in groups(data['Места'], 'Ключ'):
        reg['mes'][key] = clean({'title': first(rows, 'Название'), 'polno': first(rows, 'Описание'), 'vid': first(rows, 'Вид'),
                                 'chast': first(rows, 'Часть')})
    for key, rows in groups(data['Что передаётся'], 'Ключ'):
        g = first(rows, 'Где')
        reg['art'][key] = clean({'title': first(rows, 'Название'), 'vid': first(rows, 'Вид'),
                                 'gde': by_title[('mes', g)] if g else '', 'dogadka': first(rows, 'Догадка')})
    for key, rows in groups(data['Обязанности'], 'Ключ'):
        reg['oby'][key] = clean({'title': first(rows, 'Название'), 'polno': first(rows, 'ЗаЧтоОтвечает'),
                                 'predpolagaem': [by_title[('rol', t)] for t in many(rows, 'Предполагаем')],
                                 'kontur': first(rows, 'Контур'), 'gde': [by_title[('mes', t)] for t in many(rows, 'Где')],
                                 'istochnik': split(first(rows, 'Источник')), 'pochemu': first(rows, 'Почему'),
                                 'dogadka': first(rows, 'Догадка'), 'status': first(rows, 'Статус')}, keep=('predpolagaem',))
    procs = {}
    ptitle = {}
    for key, rows in groups(data['Процессы'], 'Ключ'):
        ptitle[first(rows, 'Название')] = key
        procs[key] = {'kontur': first(rows, 'Контур'), 'dalee_t': many(rows, 'Дальше'), 'title': first(rows, 'Название'),
                      'reglament_title': first(rows, 'ВРегламенте'), 'shagi': [], 'tolk': {}}
    inv = {v: k for k, v in OBJECTS.items()}
    for sid, rows in groups(data['Регламент'], 'IdШага'):
        p = procs[ptitle[first(rows, 'Процесс')]]
        marks = {c: first(rows, f'Рег{c}') for c in protsessy.COLUMNS}
        notes = dict(x.split(': ', 1) for x in split(first(rows, 'РегПояснения')))
        kolonki = {c: notes.get(c, MARK_TEXT[m]) for c, m in marks.items() if m}
        step = clean({'id': sid, 'stroka': first(rows, 'ОригСтрока'), 'nomer': first(rows, 'НомерШага'), 'title': first(rows, 'Шаг'),
                      'razdel': first(rows, 'Раздел'), 'primechanie': first(rows, 'РегПримечание'),
                      'v_sisteme': first(rows, 'РегВСистеме'), 'izmeneniya': first(rows, 'РегИзменения'), 'kolonki': kolonki},
                     keep=('nomer',))
        p['shagi'].append(step)
        ob = lambda k: [by_title[('oby', t)] for t in many(rows, k)]
        ro = lambda k: ['rol.' + by_title[('rol', t)] for t in many(rows, k)]
        st = {inv[r['ОбъектСтатуса']]: r['СтатусПосле'] for r in rows if r.get('ОбъектСтатуса')}
        p['tolk'][sid] = clean({'delaet': by_title[('oby', first(rows, 'Делает'))], 'stavit': ob('Ставит'), 'prinimaet': ob('Принимает'),
                                'peredaet': [by_title[('art', t)] for t in many(rows, 'Передаёт')],
                                'otslezhivaet': ob('Отслеживает'), 'soglasuyut': ob('Согласуют'), 'pomogayut': ro('Помогают'),
                                'uznayut': ro('Узнают'), 'mesto': [by_title[('mes', t)] for t in many(rows, 'Место')],
                                'status': st, 'dogadka': first(rows, 'Догадка')})
    for p in procs.values():
        p['dalee'] = [ptitle[t] for t in p.pop('dalee_t')]
    papki = [clean({'stroka': int(first(rows, 'ОригСтрока')), 'uroven': int(first(rows, 'Уровень')), 'papka': first(rows, 'Папка'),
                    'sozdaet': first(rows, 'Создаёт'), 'napolnyaet': first(rows, 'Наполняет'), 'primechanie': first(rows, 'Примечание')})
             for _, rows in groups(data['Папки'], 'Ключ')]
    return reg, procs, papki


def split(text):
    return [x.strip() for x in text.split(';') if x.strip()] if text else []


def clean(d, keep=()):
    return {k: v for k, v in d.items() if v not in ('', [], {}, None) or k in keep}


# ---------- сверка ----------

def norm_repo(reg, procs, papki):
    """Данные репозитория в том же виде, что даёт read_book: чтобы сравнить."""
    out_reg = {}
    fields = {'rol': ('title', 'polno', 'kontur', 'kolonka', 'nazvaniya', 'istochnik', 'pochemu', 'dogadka', 'status'),
              'mes': ('title', 'polno', 'vid', 'chast'),
              'art': ('title', 'vid', 'gde', 'dogadka'),
              'oby': ('title', 'polno', 'predpolagaem', 'kontur', 'gde', 'istochnik', 'pochemu', 'dogadka', 'status')}
    for kind, keys in fields.items():
        out_reg[kind] = {}
        for k, v in reg[kind].items():
            lists = ('kolonka', 'nazvaniya', 'istochnik', 'predpolagaem', 'gde') if kind != 'art' else ()   # у предмета gde — одно место
            d = {f: (protsessy.as_list(v.get(f)) if f in lists
                     else str(v.get(f, '')) if v.get(f) is not None else '') for f in keys}
            out_reg[kind][k] = clean(d, keep=('predpolagaem',) if kind == 'oby' else ())
    out_procs = {}
    for p in procs:
        shagi = []
        for s in p['shagi']:
            kol = {c: protsessy.mark(v) for c, v in (s.get('kolonki') or {}).items()}
            shagi.append(clean({'id': s['id'], 'stroka': str(s.get('stroka', '')), 'nomer': s['nomer'], 'title': s['title'],
                                'razdel': s.get('razdel', ''), 'primechanie': s.get('primechanie', ''),
                                'v_sisteme': s.get('v_sisteme', ''), 'izmeneniya': s.get('izmeneniya', ''), 'kolonki': kol},
                               keep=('nomer',)))
        tolk = {sid: clean({k: (v if k in ('delaet', 'dogadka', 'status') else protsessy.as_list(v)) for k, v in t.items()})
                for sid, t in p['tolk'].items()}
        out_procs[p['slug']] = {'kontur': p['kontur'], 'dalee': p['dalee'], 'title': title_of(p),
                                'reglament_title': p['reglament_title'], 'shagi': shagi, 'tolk': tolk}
    return out_reg, out_procs, [clean(f) for f in papki]


def norm_book(reg, procs, papki):
    for p in procs.values():
        for s in p['shagi']:
            s['kolonki'] = {c: protsessy.mark(v) for c, v in s.get('kolonki', {}).items()}
            if 'kolonki' in s and not s['kolonki']:
                del s['kolonki']
    return reg, procs, papki


def diff(a, b, path=''):
    """Различия двух структур: список строк «где: было → стало»."""
    out = []
    if isinstance(a, dict) and isinstance(b, dict):
        for k in list(a) + [k for k in b if k not in a]:
            if k not in b:
                out.append(f'{path}{k}: убрано')
            elif k not in a:
                out.append(f'{path}{k}: добавлено')
            else:
                out += diff(a[k], b[k], f'{path}{k}.')
    elif isinstance(a, list) and isinstance(b, list) and all(isinstance(x, dict) and 'id' in x for x in a + b):
        out += diff({x['id']: x for x in a}, {x['id']: x for x in b}, path)
    elif a != b:
        out.append(f'{path.rstrip(".")}: {a!r} → {b!r}')
    return out


def compare(path):
    _, reg, procs, papki = repo()
    mine = norm_repo(reg, procs, papki)
    book = norm_book(*read_book(path))
    return diff(mine[0], book[0], 'reg.') + diff(mine[1], book[1], 'pro.') + diff(
        {f['stroka']: f for f in mine[2]}, {f['stroka']: f for f in book[2]}, 'papki.')


def main():
    args = sys.argv[1:]
    if '--out' in args:
        path = args[args.index('--out') + 1]
        counts = export(path)
        print(f'книга: {path}\n  строк: ' + ', '.join(f'{k} {v}' for k, v in counts.items()))
        d = compare(path)
        print('  сверка с репозиторием: ' + ('совпадает' if not d else f'различий {len(d)}'))
        for x in d[:30]:
            print('   ', x)
        sys.exit(1 if d else 0)
    if '--check' in args:
        with tempfile.TemporaryDirectory() as tmp:
            path = f'{tmp}/check.xlsx'
            export(path)
            d = compare(path)
        print('reglament_xlsx: ' + ('выгрузка и чтение совпадают с репозиторием' if not d else f'различий {len(d)}'))
        for x in d[:30]:
            print('  ', x)
        sys.exit(1 if d else 0)
    if '--diff' in args:
        d = compare(args[args.index('--diff') + 1])
        print('отличий от репозитория: ' + str(len(d)))
        for x in d:
            print('  ', x)
        return
    sys.exit(__doc__)


if __name__ == '__main__':
    main()
